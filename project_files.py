"""Validation and bounded text extraction for persistent project attachments."""

from io import BytesIO
from pathlib import Path
from zipfile import ZipFile


MAX_EXTRACTED_TEXT_CHARS = 100_000
MAX_PDF_PAGES = 100
MAX_XLSX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_XLSX_ARCHIVE_ENTRIES = 4096
MAX_XLSX_CELLS = 50_000
MAX_XLSX_ROWS = 5_000
MAX_XLSX_COLUMNS = 100
MAX_XLSX_SHEETS = 20


class _XlsxValidationError(ValueError):
    """Expected, user-facing workbook validation failure."""


_MIME_TYPES = {
    ".csv": "text/csv",
    ".md": "text/markdown",
    ".pdf": "application/pdf",
    ".txt": "text/plain",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def extract_project_file(filename: str, content: bytes) -> tuple[str, str]:
    """Return a verified MIME type and bounded searchable text for a file."""
    extension = Path(filename).suffix.lower()
    mime_type = _MIME_TYPES.get(extension)
    if not mime_type:
        raise ValueError("Formats acceptés : CSV, Excel XLSX, PDF, TXT et Markdown.")
    if not content:
        raise ValueError("Le fichier est vide.")

    if extension == ".xlsx":
        workbook = None
        try:
            with ZipFile(BytesIO(content)) as archive:
                entries = archive.infolist()
                if len(entries) > MAX_XLSX_ARCHIVE_ENTRIES:
                    raise _XlsxValidationError("Le classeur Excel contient trop d’éléments.")
                if sum(entry.file_size for entry in entries) > MAX_XLSX_UNCOMPRESSED_BYTES:
                    raise _XlsxValidationError(
                        "Le classeur Excel est trop volumineux après décompression."
                    )

            from openpyxl import load_workbook
            from openpyxl.utils import get_column_letter

            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
            lignes = []
            caracteres = 0
            cellules = 0
            cellules_avec_valeur = 0
            lignes_lues = 0
            for feuille in workbook.worksheets[:MAX_XLSX_SHEETS]:
                if feuille.sheet_state != "visible" or lignes_lues >= MAX_XLSX_ROWS:
                    continue
                colonnes = min(max(feuille.max_column or 1, 1), MAX_XLSX_COLUMNS)
                lignes_restantes = MAX_XLSX_ROWS - lignes_lues
                titre_feuille = f"[Feuille : {feuille.title}]"
                if caracteres + len(titre_feuille) + 1 > MAX_EXTRACTED_TEXT_CHARS:
                    break
                lignes.append(titre_feuille)
                caracteres += len(titre_feuille) + 1
                for numero_ligne, row in enumerate(
                    feuille.iter_rows(
                        max_row=min(max(feuille.max_row or 1, 1), lignes_restantes),
                        max_col=colonnes,
                        values_only=True,
                    ),
                    start=1,
                ):
                    lignes_lues += 1
                    cellules += len(row)
                    valeurs = [
                        f"{get_column_letter(index)}: {str(value)[:1000]}"
                        for index, value in enumerate(row, start=1)
                        if value is not None and str(value).strip()
                    ]
                    cellules_avec_valeur += len(valeurs)
                    if valeurs:
                        ligne = f"Ligne {numero_ligne} : " + " | ".join(valeurs)
                        restant = MAX_EXTRACTED_TEXT_CHARS - caracteres
                        if restant > 0:
                            lignes.append(ligne[:restant])
                            caracteres += min(len(ligne), restant) + 1
                    if (
                        cellules >= MAX_XLSX_CELLS
                        or caracteres >= MAX_EXTRACTED_TEXT_CHARS
                        or lignes_lues >= MAX_XLSX_ROWS
                    ):
                        break
                if (
                    cellules >= MAX_XLSX_CELLS
                    or caracteres >= MAX_EXTRACTED_TEXT_CHARS
                    or lignes_lues >= MAX_XLSX_ROWS
                ):
                    break
            extracted = "\n".join(lignes)
        except _XlsxValidationError:
            raise
        except Exception as exc:
            raise ValueError("Le classeur Excel est invalide ou illisible.") from exc
        finally:
            if workbook is not None:
                workbook.close()
        if not cellules_avec_valeur:
            raise ValueError("Le classeur Excel ne contient aucune donnée lisible.")
    elif extension == ".pdf":
        if not content.startswith(b"%PDF-"):
            raise ValueError("Le fichier PDF est invalide.")
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise ValueError("La lecture PDF est indisponible sur ce serveur.") from exc
        try:
            reader = PdfReader(BytesIO(content), strict=False)
            if reader.is_encrypted:
                raise ValueError("Les PDF protégés par mot de passe ne sont pas acceptés.")
            pages = reader.pages[:MAX_PDF_PAGES]
            extracted = "\n".join(page.extract_text() or "" for page in pages)
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("Le contenu de ce PDF n’a pas pu être lu.") from exc
        if not extracted.strip():
            raise ValueError("Ce PDF ne contient pas de texte exploitable.")
    else:
        if b"\x00" in content:
            raise ValueError("Le fichier texte contient des données binaires invalides.")
        try:
            extracted = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                extracted = content.decode("cp1252")
            except UnicodeDecodeError as exc:
                raise ValueError("Le fichier texte n’utilise pas un encodage compatible.") from exc

    return mime_type, extracted[:MAX_EXTRACTED_TEXT_CHARS]

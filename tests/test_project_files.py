"""Project attachment validation and extraction limits."""

from io import BytesIO
import unittest

from openpyxl import Workbook

from project_files import (
    MAX_EXTRACTED_TEXT_CHARS,
    extract_project_file,
)


class ProjectFileProcessingTests(unittest.TestCase):
    def test_text_and_csv_use_verified_types_and_decode_windows_text(self):
        mime, text = extract_project_file("sales.csv", "ville,année\n".encode("cp1252"))
        self.assertEqual(mime, "text/csv")
        self.assertIn("année", text)

        mime, text = extract_project_file("brief.txt", b"Project brief")
        self.assertEqual(mime, "text/plain")
        self.assertEqual(text, "Project brief")

    def test_rejects_binary_text_and_unsupported_formats(self):
        with self.assertRaises(ValueError):
            extract_project_file("secret.exe", b"data")
        with self.assertRaises(ValueError):
            extract_project_file("bad.txt", b"text\x00data")
        with self.assertRaisesRegex(ValueError, "PDF est invalide"):
            extract_project_file("bad.pdf", b"not a pdf")

    def test_extracted_text_is_bounded(self):
        _, text = extract_project_file("large.txt", b"a" * (MAX_EXTRACTED_TEXT_CHARS + 1))
        self.assertEqual(len(text), MAX_EXTRACTED_TEXT_CHARS)

    def test_xlsx_extracts_visible_sheets_with_sheet_and_cell_references(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Ventes Burkina"
        sheet.append(["Ville", "Ventes"])
        sheet.append(["Ouagadougou", 1250])
        hidden = workbook.create_sheet("Interne")
        hidden.sheet_state = "hidden"
        hidden.append(["Ne pas transmettre"])
        output = BytesIO()
        workbook.save(output)

        mime, text = extract_project_file("ventes.xlsx", output.getvalue())

        self.assertEqual(
            mime,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("[Feuille : Ventes Burkina]", text)
        self.assertIn("Ligne 2 : A: Ouagadougou | B: 1250", text)
        self.assertNotIn("Ne pas transmettre", text)

    def test_invalid_xlsx_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "classeur Excel est invalide"):
            extract_project_file("broken.xlsx", b"not a workbook")

    def test_empty_xlsx_is_rejected(self):
        workbook = Workbook()
        workbook.active.delete_rows(1)
        output = BytesIO()
        workbook.save(output)

        with self.assertRaisesRegex(ValueError, "aucune donnée lisible"):
            extract_project_file("empty.xlsx", output.getvalue())

    def test_xlsx_output_is_bounded(self):
        workbook = Workbook()
        workbook.active["A1"] = "x" * (MAX_EXTRACTED_TEXT_CHARS * 2)
        output = BytesIO()
        workbook.save(output)

        _, text = extract_project_file("large.xlsx", output.getvalue())

        self.assertLessEqual(len(text), MAX_EXTRACTED_TEXT_CHARS)


if __name__ == "__main__":
    unittest.main()
